"""
Document Q&A Service
=====================
Handles:
  1. Extracting text from uploaded PDF and DOCX files
  2. Building a per-session FAISS vectorstore from the extracted text
  3. Answering user questions against that document using RAG + Groq
"""

import os
import uuid
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from langchain_core.documents import Document as LcDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_groq import ChatGroq
from langchain_core.prompts import ChatPromptTemplate

from app.config import settings

# ---------------------------------------------------------------------------
# Directory for temporary per-session FAISS stores
# ---------------------------------------------------------------------------
DOC_SESSION_DIR = Path(tempfile.gettempdir()) / "internai_doc_sessions"
DOC_SESSION_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# In-memory cache: doc_session_id -> vectorstore (avoids reloading from disk)
# ---------------------------------------------------------------------------
_vectorstore_cache: dict[str, FAISS] = {}

# ---------------------------------------------------------------------------
# DOCUMENT TEXT EXTRACTION
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract full text from a PDF file using PyMuPDF."""
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract full text from a DOCX file."""
    import io
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def extract_document_text(file_bytes: bytes, content_type: str, filename: str) -> str:
    """
    Route to the correct extractor based on MIME type or file extension.
    Returns the full text content of the document.
    """
    ct = content_type.lower()
    fn = filename.lower()

    if ct == "application/pdf" or fn.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif (
        ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or fn.endswith(".docx")
    ):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported document type: {content_type}")

    if not text.strip():
        raise ValueError("No text could be extracted from the document. It may be scanned or image-only.")

    return text


# ---------------------------------------------------------------------------
# VECTORSTORE BUILDING
# ---------------------------------------------------------------------------

def build_document_vectorstore(text: str, doc_session_id: str) -> int:
    """
    Split the document text into chunks, embed them, and store in FAISS.
    Returns the number of chunks created.
    Saves the vectorstore to disk AND caches it in memory.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=600,
        chunk_overlap=80,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    docs = [LcDocument(page_content=text, metadata={"source": doc_session_id})]
    chunks = splitter.split_documents(docs)

    if not chunks:
        raise ValueError("Document could not be split into chunks.")

    embeddings = HuggingFaceEmbeddings(model_name=settings.embedding_model)
    vectorstore = FAISS.from_documents(chunks, embeddings)

    # Persist to disk
    index_path = DOC_SESSION_DIR / doc_session_id
    vectorstore.save_local(str(index_path))

    # Cache in memory
    _vectorstore_cache[doc_session_id] = vectorstore

    return len(chunks)


def get_document_vectorstore(doc_session_id: str) -> FAISS:
    """
    Load the FAISS vectorstore for a given document session.
    Uses the in-memory cache if available, otherwise loads from disk.
    """
    if doc_session_id in _vectorstore_cache:
        return _vectorstore_cache[doc_session_id]

    index_path = DOC_SESSION_DIR / doc_session_id
    if not index_path.exists():
        raise FileNotFoundError(
            f"Document session '{doc_session_id}' not found. "
            "Please upload a document first."
        )

    embeddings = HuggingFaceEmbeddings(model_name=settings.embedding_model)
    vs = FAISS.load_local(
        str(index_path), embeddings, allow_dangerous_deserialization=True
    )
    _vectorstore_cache[doc_session_id] = vs
    return vs


# ---------------------------------------------------------------------------
# DOCUMENT Q&A
# ---------------------------------------------------------------------------

DOC_QA_SYSTEM_TEMPLATE = """You are a helpful document analysis assistant. You have been given the content of a document uploaded by the user. Use the **Document Context** below to answer the user's questions accurately and thoroughly.

## Document Context
{context}

## Instructions
- Answer questions **only based on the document content** provided above.
- If the answer is not in the document, say so clearly rather than guessing.
- When generating questions from the document, create a mix of factual, conceptual, and analytical questions.
- Format your responses with clear **Markdown** (headings, bullets, numbered lists) for readability.
- Be concise but complete.
"""


def generate_document_qa_response(doc_session_id: str, user_query: str) -> str:
    """
    Answer a user question about their uploaded document using RAG.

    Args:
        doc_session_id: ID of the document session (links to FAISS store)
        user_query: The user's question

    Returns:
        The AI response as a string
    """
    # 1. Retrieve relevant document chunks
    vectorstore = get_document_vectorstore(doc_session_id)
    retriever = vectorstore.as_retriever(search_kwargs={"k": 5})
    docs = retriever.invoke(user_query)
    context_text = "\n\n---\n\n".join([doc.page_content for doc in docs])

    # 2. Build prompt
    prompt = ChatPromptTemplate.from_messages([
        ("system", DOC_QA_SYSTEM_TEMPLATE),
        ("human", "{user_query}"),
    ])

    # 3. Invoke LLM
    if not settings.groq_api_key:
        return (
            "⚠️ The AI backend is not configured. Please add a valid `GROQ_API_KEY` "
            "to the `.env` file and restart the server."
        )

    llm = ChatGroq(
        api_key=settings.groq_api_key,
        model_name=settings.groq_model,
        temperature=0.3,
    )

    chain = prompt | llm

    response = chain.invoke({
        "context": context_text,
        "user_query": user_query,
    })

    return response.content
