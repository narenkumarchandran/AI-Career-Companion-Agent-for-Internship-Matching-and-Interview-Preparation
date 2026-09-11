import os
from pathlib import Path
from pathlib import Path
from docx import Document

from langchain_core.documents import Document as LcDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS

from app.config import settings

def load_docx(file_path: str) -> list[LcDocument]:
    doc = Document(file_path)
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return [LcDocument(page_content=text, metadata={"source": file_path})]

def build_chatbot_index():
    """
    Reads app/data/product_knowledge.docx, splits it into chunks,
    embeds them, and saves a FAISS index to app/data/faiss_chatbot_index.
    """
    doc_path = os.path.join("app", "data", "product_knowledge.docx")
    if not os.path.exists(doc_path):
        raise FileNotFoundError(f"Product knowledge document not found at {doc_path}")

    print(f"Loading document from {doc_path}...")
    docs = load_docx(doc_path)
    
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50,
        separators=["\n\n", "\n", ".", " ", ""]
    )
    chunks = text_splitter.split_documents(docs)
    
    print(f"Split document into {len(chunks)} chunks.")
    print(f"Initializing embedding model ({settings.embedding_model})...")
    
    embeddings = HuggingFaceEmbeddings(model_name=settings.embedding_model)
    
    print("Building FAISS index...")
    vectorstore = FAISS.from_documents(chunks, embeddings)
    
    index_dir = os.path.join("app", "data", "faiss_chatbot_index")
    vectorstore.save_local(index_dir)
    print(f"FAISS index saved to {index_dir}")

def get_chatbot_vectorstore():
    """
    Loads the existing FAISS index for the chatbot.
    """
    index_dir = os.path.join("app", "data", "faiss_chatbot_index")
    if not os.path.exists(index_dir):
        raise FileNotFoundError(f"Chatbot index not found at {index_dir}. Did you run build_chatbot_index.py?")
    
    embeddings = HuggingFaceEmbeddings(model_name=settings.embedding_model)
    vectorstore = FAISS.load_local(index_dir, embeddings, allow_dangerous_deserialization=True)
    return vectorstore
