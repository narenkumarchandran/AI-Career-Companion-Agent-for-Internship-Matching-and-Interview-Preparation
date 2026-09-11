import sys
from docx import Document

def create_product_knowledge_doc(output_path):
    doc = Document()
    
    doc.add_heading('Internship Assistant - Product Knowledge', 0)
    
    doc.add_heading('1. Product Overview', level=1)
    doc.add_paragraph('Product Name: Internship Assistant')
    doc.add_paragraph('Product Description: The Internship Assistant is an AI-powered platform designed to match students with the most suitable internship opportunities based on their resumes.')
    doc.add_paragraph('Problem the product solves: It eliminates the manual effort of reading through hundreds of job postings to find a good fit. By leveraging semantic search, it helps students discover roles that align with their skills and experience even if the exact keywords do not match.')
    doc.add_paragraph('Target users: University students and early-career professionals looking for internships.')
    doc.add_paragraph('Main objectives of the product: To automate resume parsing and provide highly accurate, semantically relevant internship recommendations.')
    
    doc.add_heading('2. Product Features and Functionalities', level=1)
    doc.add_heading('Feature: Resume Upload and Parsing', level=2)
    doc.add_paragraph('Purpose: To extract structured information from a student\'s resume.')
    doc.add_paragraph('How it works: Users upload their resume (PDF/DOCX). The system parses it using regular expressions and a fallback Large Language Model (LLM) to extract skills, education, and experience.')
    doc.add_paragraph('How users can use it: Navigate to the upload section and select a resume file.')
    doc.add_paragraph('Expected output or result: A structured profile containing extracted skills, work experience, education, and contact details.')
    
    doc.add_heading('Feature: Internship Matching (Semantic Search)', level=2)
    doc.add_paragraph('Purpose: To recommend the most relevant internships for the user.')
    doc.add_paragraph('How it works: The system combines the user\'s extracted skills, experience, and education into a single query text. This text is converted into an embedding using sentence-transformers. The embedding is then compared against a FAISS vector database containing pre-embedded internship postings. The top matches are returned and scored based on semantic similarity, skill overlap, education, and location.')
    doc.add_paragraph('How users can use it: After uploading a resume, the user clicks the "Find Matches" button.')
    doc.add_paragraph('Expected output or result: A ranked list of internship postings with match scores and a breakdown of why they matched.')
    
    doc.add_heading('3. How to Use the Product', level=1)
    doc.add_paragraph('1. User logs into the application using their credentials (or creates a new account).')
    doc.add_paragraph('2. User navigates to the Dashboard.')
    doc.add_paragraph('3. User uploads their resume document (PDF format).')
    doc.add_paragraph('4. The system parses the resume and displays the extracted profile data for review.')
    doc.add_paragraph('5. User navigates to the "Matches" section to see recommended internships.')
    doc.add_paragraph('6. The user can interact with the InternAI Chatbot to ask questions about the product and how it works.')
    
    doc.add_heading('4. Product Architecture', level=1)
    doc.add_paragraph('The high-level architecture consists of the following components:')
    doc.add_paragraph('Frontend: React-based Single Page Application (SPA) providing the user interface for authentication, resume upload, and viewing matches.')
    doc.add_paragraph('Backend: FastAPI application built with Python. It handles business logic, API requests, authentication (JWT), and database interactions.')
    doc.add_paragraph('LLM/AI Model: LangChain and Groq API are used for intelligent fallback parsing of unstructured resume text, and for powering the InternAI chatbot responses.')
    doc.add_paragraph('Vector Database: FAISS (Facebook AI Similarity Search) stores document embeddings of internship postings for fast Retrieval-Augmented Generation (RAG) and semantic matching.')
    doc.add_paragraph('Database: PostgreSQL stores user accounts, authentication tokens, parsed resume data, and chatbot conversation history.')
    doc.add_paragraph('Flow: User -> React Frontend -> FastAPI Backend -> LLM (Groq) & Vector DB (FAISS) -> PostgreSQL DB -> Response to User.')
    
    doc.add_heading('5. Technology Stack', level=1)
    doc.add_paragraph('Frontend: React, TypeScript, TailwindCSS, Vite. Selected for fast development, type safety, and modern UI capabilities.')
    doc.add_paragraph('Backend: Python, FastAPI. Selected for high performance, native async support, and excellent integration with machine learning ecosystems.')
    doc.add_paragraph('AI and LLM: LangChain, Groq API (ChatGroq), HuggingFace sentence-transformers. Selected for powerful, flexible text generation and local, cost-effective embeddings.')
    doc.add_paragraph('Database: PostgreSQL via SQLAlchemy ORM. Selected for robust relational data storage and schema management.')
    doc.add_paragraph('Vector Database: FAISS. Selected for lightweight, in-memory vector similarity search that does not require a separate external service infrastructure.')
    
    doc.save(output_path)
    print(f"Product Knowledge Document saved to {output_path}")

if __name__ == '__main__':
    create_product_knowledge_doc('app/data/product_knowledge.docx')
